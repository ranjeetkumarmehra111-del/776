import openpyxl
# import Excel file
def load_file():
    wb = openpyxl.load_workbook("12600695.xlsx",data_only=True)
    return wb["Daily Log"]
# here data is reading from Excel
def read_data(ws):

    sleep = []
    fitness = []
    study = []
    coding = []
    classes = []
    tracked = []
    free_time = []
    feeling = []
    satisfaction = []
    energy = []

    for row in range(6, ws.max_row + 1):

        if ws.cell(row, 2).value is not None:
            sleep.append(ws.cell(row, 2).value)

        if ws.cell(row, 3).value is not None:
            fitness.append(ws.cell(row, 3).value)

        if ws.cell(row, 4).value is not None:
            study.append(ws.cell(row, 4).value)

        if ws.cell(row, 5).value is not None:
            coding.append(ws.cell(row, 5).value)

        if ws.cell(row, 6).value is not None:
            classes.append(ws.cell(row, 6).value)

        if ws.cell(row, 9).value is not None:
            tracked.append(ws.cell(row, 9).value)

        if ws.cell(row, 10).value is not None:
            free_time.append(ws.cell(row, 10).value)

        if ws.cell(row, 11).value is not None:
            feeling.append(ws.cell(row, 11).value)

        if ws.cell(row, 12).value is not None:
            satisfaction.append(ws.cell(row, 12).value)

        if ws.cell(row, 13).value is not None:
            energy.append(ws.cell(row, 13).value)

    return sleep, fitness, study, coding, classes, tracked, free_time, feeling, satisfaction, energy


# Convert feeling into numbers
def feeling_score(feeling):
    score = {
        "Excellent": 5,
        "Good": 4,
        "Neutral": 3,
        "Average": 2,
        "Stressed": 1
    }
    result = []
    for x in feeling:
        x = str(x).strip().title()
        if x in score:
            result.append(score[x])

    return result

# Convert satisfaction into numbers
def satisfaction_score(satisfaction):
    score = {
        "Very Satisfied": 5,
        "Satisfied": 4,
        "Neutral": 3,
        "Unsatisfied": 2,
        "Very Unsatisfied": 1
    }

    result = []
    for x in satisfaction:
        x = str(x).strip().title()

        if x in score:
            result.append(score[x])

    return result


# Convert energy into numbers
def energy_score(energy):

    score = {
        "High": 3,
        "Medium": 2,
        "Low": 1
    }

    result = []

    for x in energy:
        x = str(x).strip().title()
        if x in score:
            result.append(score[x])
    return result


# Calculate project results
def calculate(sleep, fitness, study, coding, classes,tracked, free_time, feeling, satisfaction, energy):
    days = len(coding)
    if days == 0:
        return 0
    tpi = sum(coding) / days
    aai = (sum(study) + sum(classes)) / days
    phai = sum(fitness) / days

    sri = sum(sleep) / days

    abi = sum(free_time) / days

    tui = sum(tracked) / days

    ei = (sum(feeling) + sum(satisfaction) + sum(energy)) / (3 * days)

    return tpi, aai, phai, sri, abi, tui, ei


# Main function
def main():
    expected_days = 32

    # for loading file
    ws = load_file()

    # Read Excel data
    sleep, fitness, study, coding, classes, tracked, free_time, feeling, satisfaction, energy = read_data(ws)

    # Convert dictionary values
    feeling = feeling_score(feeling)
    satisfaction = satisfaction_score(satisfaction)
    energy = energy_score(energy)

    # Calculate indices
    tpi, aai, phai, sri, abi, tui, ei = calculate(sleep,fitness,study,coding,classes,tracked,free_time,feeling,satisfaction,energy)

    # Number of days
    valid_days = len(coding)

    # Data Continuity Index
    dci = (valid_days * 100) / expected_days

    # Personal Activity Index
    pai = (0.15 * tpi +0.20 * aai +0.15 * phai +0.20 * sri +0.15 * tui +0.10 * ei+0.05 * dci)

    # Display result
    print("*" * 50)
    print("          MY DATA, MY STORY")
    print("*" * 50)

    print("Expected Days:", expected_days)
    print("Recorded Days:", valid_days)
    print("Missing Days:", expected_days - valid_days)

    print()
    print("Activity Results")
    print("-" * 50)

    print("Tech Productivity:", round(tpi, 2))
    print("Academic Activity:", round(aai, 2))
    print("Physical Activity:", round(phai, 2))
    print("Sleep & Recovery:", round(sri, 2))
    print("Activity Balance:", round(abi, 2))
    print("Time Utilization:", round(tui, 2))
    print("Experience:", round(ei, 2))
    print("Data Continuity:", round(dci, 2))

    print()
    print("Personal Activity Index:", round(pai, 2))

    print()
    print("Project completed successfully")


# Start program
main()